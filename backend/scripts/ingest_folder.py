"""Ingest an entire folder of course material (recursively) into ChromaDB
under one course_id. Supports .pdf, .pptx, and .ipynb files. Other file
types (.ppt, .docx, images, etc.) are skipped with a warning — convert them
first if you need them included (e.g. re-save .ppt as .pptx in PowerPoint).

Usage (run from the `backend/` directory so the `app` package resolves):
    python scripts/ingest_folder.py path/to/folder COURSE_ID

Example:
    python scripts/ingest_folder.py data/unit1 deep-learning
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.tools.vector_store import ingest_document  # noqa: E402

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".ipynb"}


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
        else:
            print(f"    [warn] page {i + 1} had no extractable text (scanned image?)")
    return "\n\n".join(parts)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        slide_text.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            slide_text.append(cell.text)
        # Speaker notes, if present
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                slide_text.append(f"[Notes] {notes}")
        if slide_text:
            parts.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_text))
    return "\n\n".join(parts)


def extract_ipynb(path: Path) -> str:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    parts = []
    for cell in notebook.get("cells", []):
        source = "".join(cell.get("source", []))
        if not source.strip():
            continue
        cell_type = cell.get("cell_type", "code")
        if cell_type == "markdown":
            parts.append(source)
        elif cell_type == "code":
            # Only markdown + code are kept — printed outputs (training logs,
            # loss/epoch progress, etc.) are noise for RAG and bloat ingestion.
            parts.append(f"```python\n{source}\n```")
    return "\n\n".join(parts)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    elif suffix == ".pptx":
        return extract_pptx(path)
    elif suffix == ".ipynb":
        return extract_ipynb(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def main():
    if len(sys.argv) != 3:
        print("Usage: python scripts/ingest_folder.py <folder> <course_id>")
        sys.exit(1)

    folder = Path(sys.argv[1]).resolve()
    course_id = sys.argv[2]

    if not folder.is_dir():
        print(f"Not a folder: {folder}")
        sys.exit(1)

    all_files = sorted(p for p in folder.rglob("*") if p.is_file())
    total_chunks = 0
    ingested = []
    skipped = []

    for path in all_files:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            skipped.append(path)
            continue

        rel = path.relative_to(folder)
        print(f"Processing {rel} ...")
        try:
            text = extract_text(path)
        except Exception as e:
            print(f"  [error] failed to extract text: {e}")
            skipped.append(path)
            continue

        if not text.strip():
            print("  [warn] no text extracted — skipping (scanned image? needs OCR)")
            skipped.append(path)
            continue

        n_chunks = ingest_document(
            course_id=course_id,
            text=text,
            source_name=str(rel),
        )
        total_chunks += n_chunks
        ingested.append(rel)
        print(f"  -> {n_chunks} chunks stored")

    print("\n=== Summary ===")
    print(f"Course ID: {course_id}")
    print(f"Ingested: {len(ingested)} file(s), {total_chunks} chunks total")
    for rel in ingested:
        print(f"  OK   {rel}")
    if skipped:
        print(f"Skipped: {len(skipped)} file(s) (unsupported type, no text, or error)")
        for path in skipped:
            print(f"  SKIP {path.relative_to(folder)}")
    print(f"\nTry: POST /api/v1/tutor/ask with course_id='{course_id}'")


if __name__ == "__main__":
    main()
